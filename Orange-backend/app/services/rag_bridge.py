"""
RAG 模块桥接层
封装对独立 RAG 模块 RAGApi 的调用
当 RAG 模块作为独立包安装后，通过此模块统一调用
"""

from typing import Optional

from app.schemas.rag_dto import DocumentEntity, DocumentIngestResult, DocumentSearchResult
from app.utils.logger import logger


class RAGApiBridge:
    """
    RAG 模块桥接器

    当独立 RAG 模块（citrus_agent）可用时，直接调用 RAGApi；
    当不可用时，回退到内置的简化实现（基于 core/ 下的 Qdrant/Embedding/Chunking）
    """

    def __init__(self):
        self._rag_api = None
        self._use_builtin = False
        self._init_rag_api()

    def _init_rag_api(self):
        """尝试初始化 RAGApi，失败则回退到内置实现"""
        try:
            from src.citrus_agent.rag.rag_api import RAGApi
            self._rag_api = RAGApi()
            logger.info("RAGApi（独立模块）初始化成功")
        except ImportError:
            logger.info("独立 RAG 模块不可用，使用内置实现")
            self._use_builtin = True

    def ingest_document(self, document: DocumentEntity) -> DocumentIngestResult:
        """
        文档入库

        Args:
            document: 文档实体

        Returns:
            入库结果
        """
        if self._use_builtin:
            return self._builtin_ingest(document)

        return self._rag_api.ingest_document(document)

    def delete_document_vectors(self, document_id: int) -> None:
        """按 document_id 删除向量"""
        if self._use_builtin:
            return self._builtin_delete_document(document_id)

        self._rag_api.delete_document_vectors(document_id=document_id)

    def delete_knowledge_base_vectors(self, kb_id: int) -> None:
        """按 kb_id 删除向量"""
        if self._use_builtin:
            return self._builtin_delete_kb(kb_id)

        self._rag_api.delete_knowledge_base_vectors(kb_id=kb_id)

    def search(
        self,
        query_text: str,
        top_k: int = 5,
        kb_id: Optional[int] = None,
    ) -> DocumentSearchResult:
        """
        知识库检索

        Args:
            query_text: 查询文本
            top_k: 返回数量
            kb_id: 知识库ID过滤

        Returns:
            检索结果
        """
        if self._use_builtin:
            return self._builtin_search(query_text, top_k, kb_id)

        return self._rag_api.search(query_text=query_text, top_k=top_k, kb_id=kb_id)

    # ==================== 内置回退实现 ====================

    def _builtin_ingest(self, document: DocumentEntity) -> DocumentIngestResult:
        """内置入库实现（同步调用，简化版）"""
        import asyncio
        return asyncio.get_event_loop().run_until_complete(
            self._async_builtin_ingest(document)
        )

    async def _async_builtin_ingest(self, document: DocumentEntity) -> DocumentIngestResult:
        """异步内置入库"""
        try:
            from app.core.qdrant import get_qdrant_manager
            from app.core.embedding import get_embedding_client
            from app.core.chunking import DocumentChunker
            from qdrant_client.models import PointStruct
            import uuid

            qdrant = get_qdrant_manager()
            embedding = get_embedding_client()
            chunker = DocumentChunker()

            # 读取文件
            from app.services.knowledge_service import _get_mime_type
            ext = document.stored_path.rsplit(".", 1)[-1].lower() if "." in document.stored_path else ""
            content = await self._read_file_builtin(document.stored_path, ext)

            if not content:
                return DocumentIngestResult(
                    document_id=document.id,
                    kb_id=document.kb_id,
                    success=False,
                    error_message="文件内容为空或解析失败",
                )

            # 切片
            chunks = chunker.chunk_text(
                text=content,
                metadata={
                    "document_id": document.id,
                    "kb_id": document.kb_id,
                    "file_name": document.original_filename,
                },
            )

            # 向量化
            texts = [chunk.content for chunk in chunks]
            vectors = await embedding.aembed_documents(texts)

            # 先删除旧向量
            try:
                qdrant.delete_points_by_filter(
                    field="document_id", value=document.id
                )
            except Exception:
                pass

            # 写入新向量
            points = []
            for i, (chunk, vector) in enumerate(zip(chunks, vectors)):
                chunk_id = f"{document.kb_id}_{document.id}_{i}"
                points.append(
                    PointStruct(
                        id=str(uuid.uuid4()),
                        vector=vector,
                        payload={
                            "chunk_id": chunk_id,
                            "kb_id": document.kb_id,
                            "document_id": document.id,
                            "file_name": document.original_filename,
                            "page": None,
                            "chunk_index": i,
                            "text": chunk.content,
                        },
                    )
                )

            if points:
                qdrant.upsert_points(points)

            return DocumentIngestResult(
                document_id=document.id,
                kb_id=document.kb_id,
                success=True,
                chunk_count=len(chunks),
            )

        except Exception as e:
            logger.error(f"内置入库失败: {e}")
            return DocumentIngestResult(
                document_id=document.id,
                kb_id=document.kb_id,
                success=False,
                error_message=str(e),
            )

    def _builtin_delete_document(self, document_id: int) -> None:
        """内置删除文档向量"""
        try:
            from app.core.qdrant import get_qdrant_manager
            qdrant = get_qdrant_manager()
            # 按 payload 中的 document_id 过滤删除
            from qdrant_client.models import Filter, FieldCondition, MatchValue
            qdrant.client.delete(
                collection_name=qdrant.collection_name,
                points_selector=Filter(
                    must=[FieldCondition(key="document_id", match=MatchValue(value=document_id))]
                ),
            )
            logger.info(f"已删除 document_id={document_id} 的向量")
        except Exception as e:
            logger.warning(f"内置删除文档向量失败: {e}")

    def _builtin_delete_kb(self, kb_id: int) -> None:
        """内置删除知识库向量"""
        try:
            from app.core.qdrant import get_qdrant_manager
            from qdrant_client.models import Filter, FieldCondition, MatchValue
            qdrant = get_qdrant_manager()
            qdrant.client.delete(
                collection_name=qdrant.collection_name,
                points_selector=Filter(
                    must=[FieldCondition(key="kb_id", match=MatchValue(value=kb_id))]
                ),
            )
            logger.info(f"已删除 kb_id={kb_id} 的向量")
        except Exception as e:
            logger.warning(f"内置删除知识库向量失败: {e}")

    def _builtin_search(
        self, query_text: str, top_k: int, kb_id: Optional[int]
    ) -> DocumentSearchResult:
        """内置检索实现"""
        import asyncio
        return asyncio.get_event_loop().run_until_complete(
            self._async_builtin_search(query_text, top_k, kb_id)
        )

    async def _async_builtin_search(
        self, query_text: str, top_k: int, kb_id: Optional[int]
    ) -> DocumentSearchResult:
        """异步内置检索"""
        try:
            from app.core.qdrant import get_qdrant_manager
            from app.core.embedding import get_embedding_client
            from app.schemas.rag_dto import SearchResult

            qdrant = get_qdrant_manager()
            embedding = get_embedding_client()

            query_vector = await embedding.aembed_query(query_text)

            filters = None
            if kb_id is not None:
                from qdrant_client.models import Filter, FieldCondition, MatchValue
                filters = Filter(
                    must=[FieldCondition(key="kb_id", match=MatchValue(value=kb_id))]
                )

            results = qdrant.search(
                query_vector=query_vector,
                top_k=top_k,
                filters=filters,
            )

            chunks = [
                SearchResult(
                    chunk_id=r.get("payload", {}).get("chunk_id", str(r.get("id", ""))),
                    content=r.get("payload", {}).get("text", r.get("payload", {}).get("content", "")),
                    score=r.get("score", 0),
                    payload=r.get("payload", {}),
                )
                for r in results
            ]

            max_score = max((c.score for c in chunks), default=0.0)

            return DocumentSearchResult(
                kb_id=kb_id,
                query_text=query_text,
                top_k=top_k,
                score=max_score,
                chunks=chunks,
            )

        except Exception as e:
            logger.error(f"内置检索失败: {e}")
            return DocumentSearchResult(
                kb_id=kb_id,
                query_text=query_text,
                top_k=top_k,
                score=0.0,
                chunks=[],
            )

    @staticmethod
    async def _read_file_builtin(file_path: str, ext: str) -> str:
        """内置文件读取"""
        from pathlib import Path

        if ext in ("txt", "md"):
            with open(file_path, "r", encoding="utf-8") as f:
                return f.read()
        elif ext == "pdf":
            try:
                import fitz
                doc = fitz.open(file_path)
                text = "".join(page.get_text() for page in doc)
                doc.close()
                return text
            except ImportError:
                return ""
        elif ext in ("docx", "doc"):
            try:
                from docx import Document as DocxDocument
                doc = DocxDocument(file_path)
                return "\n".join(p.text for p in doc.paragraphs)
            except ImportError:
                return ""
        elif ext in ("xlsx", "xls"):
            try:
                import openpyxl
                wb = openpyxl.load_workbook(file_path)
                parts = []
                for sheet in wb.worksheets:
                    for row in sheet.iter_rows(values_only=True):
                        parts.append(" | ".join(str(c) for c in row if c))
                return "\n".join(parts)
            except ImportError:
                return ""
        elif ext in ("csv", "json"):
            with open(file_path, "r", encoding="utf-8") as f:
                return f.read()
        return ""


# 全局单例
_rag_api_bridge: RAGApiBridge | None = None


def get_rag_api() -> RAGApiBridge:
    """获取 RAGApi 桥接器单例"""
    global _rag_api_bridge
    if _rag_api_bridge is None:
        _rag_api_bridge = RAGApiBridge()
    return _rag_api_bridge
